from datetime import datetime

from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import AttendanceRecord, Employee
from .serializers import AttendanceRecordSerializer, EmployeeSerializer
from .utils import compute_night_diff_minutes, compute_time_in_status


def parse_shift_time(value, field_name):
    if not value:
        return None
    try:
        return datetime.strptime(value, "%H:%M").time()
    except ValueError as error:
        raise ValueError(f"{field_name} must use HH:MM format.") from error


def get_employee(request):
    """The signed-in person's profile. Identity comes from the session, never the request body."""
    employee = getattr(request.user, "employee", None)
    if employee is None:
        employee = Employee.objects.create(user=request.user, name=request.user.username)
    return employee


def get_open_record(employee):
    """The session this person timed in to but has not timed out of, regardless of date."""
    return (
        employee.records.filter(first_in__isnull=False, last_out__isnull=True)
        .order_by("-date")
        .first()
    )


def filter_by_month(queryset, month):
    """Return (queryset, error_response); exactly one of the two is set."""
    if not month:
        return queryset, None
    try:
        month_start = datetime.strptime(month, "%Y-%m").date()
    except ValueError:
        return None, Response(
            {"detail": "Month must use YYYY-MM format."}, status=status.HTTP_400_BAD_REQUEST
        )
    return queryset.filter(date__year=month_start.year, date__month=month_start.month), None


@api_view(["GET", "PUT"])
@permission_classes([IsAuthenticated])
def profile(request):
    """Read the signed-in person's profile, or save their shift."""
    employee = get_employee(request)

    if request.method == "PUT":
        name = (request.data.get("name") or "").strip()
        try:
            start = parse_shift_time(request.data.get("start"), "Shift start")
            end = parse_shift_time(request.data.get("end"), "Shift end")
        except ValueError as parse_error:
            return Response({"detail": str(parse_error)}, status=status.HTTP_400_BAD_REQUEST)

        updated = []
        if name:
            employee.name = name[:64]
            updated.append("name")
        if start or end:
            if not start or not end:
                return Response(
                    {"detail": "Both start and end are required."},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            employee.shift_start = start
            employee.shift_end = end
            updated += ["shift_start", "shift_end"]
        if updated:
            employee.save(update_fields=updated)

    return Response(EmployeeSerializer(employee).data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def records(request):
    employee = get_employee(request)

    queryset, error = filter_by_month(employee.records.all(), request.query_params.get("month"))
    if error:
        return error

    totals = queryset.aggregate(work_minutes=Sum("work_minutes"), night_diff_minutes=Sum("night_diff_minutes"))
    open_record = get_open_record(employee)
    return Response(
        {
            "records": AttendanceRecordSerializer(queryset, many=True).data,
            "open_record": AttendanceRecordSerializer(open_record).data if open_record else None,
            "totals": {
                "work_minutes": totals["work_minutes"] or 0,
                "night_diff_minutes": totals["night_diff_minutes"] or 0,
            },
        }
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def clock_in(request):
    employee = get_employee(request)

    work_location = request.data.get("work_location")
    if work_location not in AttendanceRecord.WorkLocation.values:
        return Response(
            {"detail": "Choose either RTO (return to office) or WFH (work from home)."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    try:
        shift_start = parse_shift_time(request.data.get("shift_start"), "Shift start")
        shift_end = parse_shift_time(request.data.get("shift_end"), "Shift end")
    except ValueError as parse_error:
        return Response({"detail": str(parse_error)}, status=status.HTTP_400_BAD_REQUEST)

    open_record = get_open_record(employee)
    if open_record:
        return Response(
            {
                "detail": (
                    f"You already timed in on {open_record.date} at "
                    f"{timezone.localtime(open_record.first_in):%I:%M %p}. Time out first."
                )
            },
            status=status.HTTP_400_BAD_REQUEST,
        )

    if shift_start and shift_end:
        # Remember the shift the user clocked in with so it persists across devices.
        employee.shift_start = shift_start
        employee.shift_end = shift_end
        employee.save(update_fields=["shift_start", "shift_end"])
    else:
        shift_start = shift_start or employee.shift_start
        shift_end = shift_end or employee.shift_end

    if not shift_start or not shift_end:
        return Response(
            {"detail": "Set your shift before timing in."}, status=status.HTTP_400_BAD_REQUEST
        )

    today = timezone.localdate()
    now = timezone.now()
    time_in_status = compute_time_in_status(now, shift_start)
    late_reason = str(request.data.get("late_reason") or "").strip()
    if time_in_status == "Late arrival" and not late_reason:
        return Response(
            {"detail": "Explain why you are late before timing in."},
            status=status.HTTP_400_BAD_REQUEST,
        )
    defaults = {
        "work_location": work_location,
        "first_in": now,
        "shift_start": shift_start,
        "shift_end": shift_end,
        "time_in_status": time_in_status,
        "late_reason": late_reason if time_in_status == "Late arrival" else "",
    }
    record, created = AttendanceRecord.objects.get_or_create(
        employee=employee, date=today, defaults=defaults
    )
    if not created:
        record.work_location = work_location
        record.first_in = now
        record.shift_start = shift_start
        record.shift_end = shift_end
        record.time_in_status = time_in_status
        record.late_reason = late_reason if time_in_status == "Late arrival" else ""
        record.save(update_fields=["work_location", "first_in", "shift_start", "shift_end", "time_in_status", "late_reason"])
    return Response(
        AttendanceRecordSerializer(record).data,
        status=status.HTTP_201_CREATED if created else status.HTTP_200_OK,
    )


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def clock_out(request):
    employee = get_employee(request)

    # Match the currently open session rather than strictly "today's" record,
    # so overnight shifts that cross midnight are closed out correctly.
    record = get_open_record(employee)
    if not record:
        return Response({"detail": "Time in before you can time out."}, status=status.HTTP_400_BAD_REQUEST)
    now = timezone.now()
    elapsed_minutes = int((now - record.first_in).total_seconds() // 60)
    record.last_out = now
    record.work_minutes = max(0, elapsed_minutes)
    record.night_diff_minutes = compute_night_diff_minutes(record.first_in, now)
    record.save(update_fields=["last_out", "work_minutes", "night_diff_minutes"])
    return Response(AttendanceRecordSerializer(record).data)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def export_excel(request):
    employee = get_employee(request)

    month = request.query_params.get("month")
    queryset, error = filter_by_month(employee.records.all(), month)
    if error:
        return error

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"

    headers = ["Date", "Work Location", "Shift Start", "Shift End", "Time In", "Time Out", "Time In Label", "Late Reason", "Hours Worked", "Night Diff Hours"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for record in queryset.order_by("date"):
        shift_start = record.shift_start.strftime("%I:%M %p") if record.shift_start else ""
        shift_end = record.shift_end.strftime("%I:%M %p") if record.shift_end else ""
        time_in = timezone.localtime(record.first_in).strftime("%I:%M %p") if record.first_in else ""
        time_out = timezone.localtime(record.last_out).strftime("%I:%M %p") if record.last_out else ""
        hours_worked = round(record.work_minutes / 60, 2) if record.work_minutes else 0
        night_diff_hours = round(record.night_diff_minutes / 60, 2) if record.night_diff_minutes else 0
        sheet.append([
            record.date.strftime("%Y-%m-%d"),
            record.work_location,
            shift_start,
            shift_end,
            time_in,
            time_out,
            record.time_in_status,
            record.late_reason,
            hours_worked,
            night_diff_hours,
        ])

    for index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(index)
        sheet.column_dimensions[column_letter].width = max(14, len(header) + 4)

    filename = f"attendance_{employee.user.username}_{month or 'all'}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response
