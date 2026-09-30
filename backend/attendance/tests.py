import io
from datetime import datetime, time, timedelta
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from .models import AttendanceRecord, Employee

User = get_user_model()


def make_person(username, password="str0ng-pass-phrase", **shift):
    user = User.objects.create_user(username=username, password=password)
    employee = Employee.objects.create(user=user, name=username.title(), **shift)
    return user, employee


class AttendanceApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user, self.employee = make_person(
            "hanz", shift_start=time(9, 0), shift_end=time(18, 0)
        )
        self.client.force_authenticate(user=self.user)

    def clock_in(self, **payload):
        body = {"work_location": "RTO"}
        body.update(payload)
        return self.client.post(reverse("clock-in"), body)

    def clock_out(self):
        return self.client.post(reverse("clock-out"))

    def test_records_endpoint_starts_empty(self):
        response = self.client.get(reverse("records"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["records"], [])
        self.assertEqual(response.data["totals"]["work_minutes"], 0)

    def test_clock_in_creates_todays_record(self):
        self.assertEqual(AttendanceRecord.objects.count(), 0)
        response = self.clock_in()
        self.assertEqual(response.status_code, 201)
        self.assertIsNotNone(response.data["first_in"])
        self.assertIsNone(response.data["last_out"])
        self.assertEqual(AttendanceRecord.objects.count(), 1)

    def test_clock_in_requires_work_location(self):
        response = self.client.post(reverse("clock-in"), {})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_clock_in_requires_a_configured_shift(self):
        user, _ = make_person("shiftless")
        self.client.force_authenticate(user=user)
        response = self.clock_in()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_double_clock_in_is_rejected(self):
        self.clock_in()
        response = self.clock_in()
        self.assertEqual(response.status_code, 400)

    def test_clock_in_attaches_shift_and_labels_almost_late(self):
        today = timezone.localdate()
        almost_late = timezone.make_aware(datetime.combine(today, time(8, 55)))

        with mock.patch("django.utils.timezone.now", return_value=almost_late):
            response = self.clock_in(shift_start="09:00", shift_end="18:00")

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["shift_start"], "09:00:00")
        self.assertEqual(response.data["shift_end"], "18:00:00")
        self.assertEqual(response.data["time_in_status"], "Almost late")

    def test_clock_in_labels_early_bird(self):
        today = timezone.localdate()
        early = timezone.make_aware(datetime.combine(today, time(7, 45)))

        with mock.patch("django.utils.timezone.now", return_value=early):
            response = self.clock_in(shift_start="09:00", work_location="WFH")

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["time_in_status"], "Early bird")
        self.assertEqual(response.data["work_location"], "WFH")

    def test_clock_in_labels_late_arrival(self):
        today = timezone.localdate()
        late = timezone.make_aware(datetime.combine(today, time(9, 1)))

        with mock.patch("django.utils.timezone.now", return_value=late):
            response = self.clock_in(shift_start="09:00")

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["time_in_status"], "Late arrival")

    def test_clock_in_rejects_invalid_shift_time(self):
        response = self.clock_in(shift_start="9am")
        self.assertEqual(response.status_code, 400)

    def test_clock_in_falls_back_to_saved_shift(self):
        self.client.put(reverse("profile"), {"start": "18:00", "end": "03:00"}, format="json")
        today = timezone.localdate()
        clock_in_at = timezone.make_aware(datetime.combine(today, time(17, 51)))

        with mock.patch("django.utils.timezone.now", return_value=clock_in_at):
            response = self.clock_in()

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["shift_start"], "18:00:00")
        self.assertEqual(response.data["time_in_status"], "Almost late")

    def test_records_exposes_open_overnight_session(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        AttendanceRecord.objects.create(
            employee=self.employee,
            date=yesterday,
            first_in=timezone.make_aware(datetime.combine(yesterday, time(18, 0))),
            shift_start=time(18, 0),
            shift_end=time(3, 0),
        )

        response = self.client.get(reverse("records"))
        self.assertEqual(response.status_code, 200)
        self.assertIsNotNone(response.data["open_record"])
        self.assertEqual(response.data["open_record"]["date"], yesterday.isoformat())

    def test_clock_out_closes_session_started_yesterday(self):
        yesterday = timezone.localdate() - timedelta(days=1)
        started_at = timezone.make_aware(datetime.combine(yesterday, time(18, 0)))
        AttendanceRecord.objects.create(
            employee=self.employee, date=yesterday, first_in=started_at, shift_start=time(18, 0)
        )

        with mock.patch("django.utils.timezone.now", return_value=started_at + timedelta(hours=9)):
            response = self.clock_out()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["date"], yesterday.isoformat())
        self.assertAlmostEqual(response.data["work_minutes"], 540, delta=1)

    def test_clock_out_requires_clock_in_first(self):
        response = self.clock_out()
        self.assertEqual(response.status_code, 400)

    def test_double_clock_out_is_rejected(self):
        self.clock_in()
        self.clock_out()
        response = self.clock_out()
        self.assertEqual(response.status_code, 400)

    def test_clock_out_records_last_out_and_work_minutes(self):
        self.clock_in()
        record = AttendanceRecord.objects.get(employee=self.employee, date=timezone.localdate())
        record.first_in = timezone.now() - timedelta(hours=8)
        record.save(update_fields=["first_in"])

        response = self.clock_out()
        self.assertEqual(response.status_code, 200)
        record.refresh_from_db()
        self.assertIsNotNone(record.last_out)
        self.assertAlmostEqual(record.work_minutes, 480, delta=1)

    def test_night_diff_is_computed_for_overnight_shift(self):
        self.clock_in()
        record = AttendanceRecord.objects.get(employee=self.employee, date=timezone.localdate())
        today = timezone.localdate()
        shift_start = timezone.make_aware(datetime.combine(today, time(21, 0)))
        record.first_in = shift_start
        record.save(update_fields=["first_in"])

        with mock.patch("django.utils.timezone.now", return_value=shift_start + timedelta(hours=8)):
            response = self.clock_out()

        self.assertEqual(response.status_code, 200)
        record.refresh_from_db()
        # Shift 9 PM -> 5 AM: night diff applies from 10 PM -> 5 AM = 7 hours = 420 minutes.
        self.assertAlmostEqual(record.night_diff_minutes, 420, delta=1)

    def test_day_shift_has_no_night_diff(self):
        self.clock_in()
        record = AttendanceRecord.objects.get(employee=self.employee, date=timezone.localdate())
        today = timezone.localdate()
        shift_start = timezone.make_aware(datetime.combine(today, time(8, 0)))
        record.first_in = shift_start
        record.save(update_fields=["first_in"])

        with mock.patch("django.utils.timezone.now", return_value=shift_start + timedelta(hours=8)):
            response = self.clock_out()

        self.assertEqual(response.status_code, 200)
        record.refresh_from_db()
        self.assertEqual(record.night_diff_minutes, 0)

    def test_export_returns_xlsx_with_recorded_rows(self):
        today = timezone.localdate()
        midnight = timezone.make_aware(datetime.combine(today, datetime.min.time()))
        AttendanceRecord.objects.create(
            employee=self.employee,
            date=today,
            work_location="WFH",
            shift_start=time(8, 0),
            shift_end=time(17, 0),
            time_in_status="Right on time",
            first_in=midnight + timedelta(hours=8),
            last_out=midnight + timedelta(hours=17),
            work_minutes=540,
            night_diff_minutes=60,
        )

        response = self.client.get(reverse("export-excel"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        workbook = load_workbook(filename=io.BytesIO(response.content))
        sheet = workbook.active
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(
            header,
            [
                "Date",
                "Work Location",
                "Shift Start",
                "Shift End",
                "Time In",
                "Time Out",
                "Time In Label",
                "Hours Worked",
                "Night Diff Hours",
            ],
        )
        data_row = [cell.value for cell in sheet[2]]
        self.assertEqual(data_row[0], today.strftime("%Y-%m-%d"))
        self.assertEqual(data_row[1], "WFH")
        self.assertEqual(data_row[6], "Right on time")
        self.assertEqual(data_row[7], 9.0)
        self.assertEqual(data_row[8], 1.0)


class AuthenticationRequiredTests(TestCase):
    """Every attendance endpoint must refuse anonymous callers."""

    def setUp(self):
        self.client = APIClient()
        make_person("hanz", shift_start=time(9, 0), shift_end=time(18, 0))

    def test_records_rejects_anonymous(self):
        self.assertEqual(self.client.get(reverse("records")).status_code, 403)

    def test_profile_rejects_anonymous(self):
        self.assertEqual(self.client.get(reverse("profile")).status_code, 403)

    def test_clock_in_rejects_anonymous(self):
        response = self.client.post(reverse("clock-in"), {"work_location": "RTO"})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(AttendanceRecord.objects.count(), 0)

    def test_clock_out_rejects_anonymous(self):
        self.assertEqual(self.client.post(reverse("clock-out")).status_code, 403)

    def test_export_rejects_anonymous(self):
        self.assertEqual(self.client.get(reverse("export-excel")).status_code, 403)

    def test_knowing_a_name_no_longer_grants_access(self):
        """The old hole: passing someone's name used to return their timesheet."""
        response = self.client.get(reverse("records"), {"employee": "hanz"})
        self.assertEqual(response.status_code, 403)


class RegistrationTests(TestCase):
    def setUp(self):
        self.client = APIClient()

    def test_register_creates_account_profile_and_signs_in(self):
        response = self.client.post(
            reverse("auth-register"),
            {"username": "Mara", "password": "str0ng-pass-phrase", "name": "Mara Cruz"},
        )
        self.assertEqual(response.status_code, 201)
        self.assertTrue(response.data["authenticated"])
        self.assertEqual(response.data["profile"]["username"], "mara")
        self.assertEqual(response.data["profile"]["name"], "Mara Cruz")
        self.assertFalse(response.data["profile"]["shift_configured"])

        # The registration response also signs the browser in.
        self.assertEqual(self.client.get(reverse("records")).status_code, 200)

    def test_usernames_are_case_insensitive_and_unique(self):
        self.client.post(
            reverse("auth-register"), {"username": "mara", "password": "str0ng-pass-phrase"}
        )
        response = self.client.post(
            reverse("auth-register"), {"username": "MARA", "password": "another-str0ng-pass"}
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(User.objects.filter(username="mara").count(), 1)

    def test_weak_passwords_are_rejected(self):
        response = self.client.post(
            reverse("auth-register"), {"username": "mara", "password": "123"}
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(User.objects.filter(username="mara").exists())

    def test_blank_username_is_rejected(self):
        response = self.client.post(
            reverse("auth-register"), {"username": "  ", "password": "str0ng-pass-phrase"}
        )
        self.assertEqual(response.status_code, 400)

    def test_failed_registration_leaves_no_half_built_account(self):
        self.client.post(
            reverse("auth-register"), {"username": "mara", "password": "str0ng-pass-phrase"}
        )
        self.client.post(
            reverse("auth-register"), {"username": "mara", "password": "str0ng-pass-phrase"}
        )
        self.assertEqual(Employee.objects.filter(user__username="mara").count(), 1)


class SessionTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.password = "str0ng-pass-phrase"
        make_person("hanz", password=self.password, shift_start=time(9, 0), shift_end=time(18, 0))

    def test_login_then_access_then_logout(self):
        self.assertEqual(self.client.get(reverse("records")).status_code, 403)

        response = self.client.post(
            reverse("auth-login"), {"username": "hanz", "password": self.password}
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["authenticated"])
        self.assertEqual(self.client.get(reverse("records")).status_code, 200)

        self.assertEqual(self.client.post(reverse("auth-logout")).status_code, 200)
        self.assertEqual(self.client.get(reverse("records")).status_code, 403)

    def test_login_accepts_any_casing_of_the_username(self):
        response = self.client.post(
            reverse("auth-login"), {"username": "HANZ", "password": self.password}
        )
        self.assertEqual(response.status_code, 200)

    def test_wrong_password_is_rejected(self):
        response = self.client.post(
            reverse("auth-login"), {"username": "hanz", "password": "wrong-password"}
        )
        self.assertEqual(response.status_code, 401)

    def test_unknown_and_wrong_password_look_identical(self):
        """So the endpoint cannot be used to discover which usernames exist."""
        missing = self.client.post(
            reverse("auth-login"), {"username": "ghost", "password": "wrong-password"}
        )
        wrong = self.client.post(
            reverse("auth-login"), {"username": "hanz", "password": "wrong-password"}
        )
        self.assertEqual(missing.status_code, wrong.status_code)
        self.assertEqual(missing.data["detail"], wrong.data["detail"])

    def test_session_endpoint_reports_anonymous_then_signed_in(self):
        response = self.client.get(reverse("auth-session"))
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.data["authenticated"])

        self.client.post(reverse("auth-login"), {"username": "hanz", "password": self.password})
        response = self.client.get(reverse("auth-session"))
        self.assertTrue(response.data["authenticated"])
        self.assertEqual(response.data["profile"]["username"], "hanz")

    def test_session_endpoint_sets_the_csrf_cookie(self):
        response = self.client.get(reverse("auth-session"))
        self.assertIn("csrftoken", response.cookies)

    def test_password_change_requires_the_current_password(self):
        self.client.post(reverse("auth-login"), {"username": "hanz", "password": self.password})
        response = self.client.post(
            reverse("auth-password"),
            {"current_password": "wrong-password", "new_password": "another-str0ng-pass"},
        )
        self.assertEqual(response.status_code, 400)

    def test_password_change_keeps_the_browser_signed_in(self):
        self.client.post(reverse("auth-login"), {"username": "hanz", "password": self.password})
        response = self.client.post(
            reverse("auth-password"),
            {"current_password": self.password, "new_password": "another-str0ng-pass"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.client.get(reverse("records")).status_code, 200)


class MultiUserIsolationTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.hanz_user, self.hanz = make_person(
            "hanz", shift_start=time(9, 0), shift_end=time(18, 0)
        )
        self.mara_user, self.mara = make_person(
            "mara", shift_start=time(21, 0), shift_end=time(6, 0)
        )

    def clock_in(self, user, location="RTO"):
        self.client.force_authenticate(user=user)
        return self.client.post(reverse("clock-in"), {"work_location": location})

    def test_two_people_can_time_in_on_the_same_day(self):
        self.assertEqual(self.clock_in(self.hanz_user).status_code, 201)
        self.assertEqual(self.clock_in(self.mara_user, "WFH").status_code, 201)
        self.assertEqual(AttendanceRecord.objects.count(), 2)

    def test_records_only_returns_your_own_days(self):
        self.clock_in(self.hanz_user)
        self.clock_in(self.mara_user, "WFH")

        self.client.force_authenticate(user=self.mara_user)
        response = self.client.get(reverse("records"))
        self.assertEqual(len(response.data["records"]), 1)
        self.assertEqual(response.data["records"][0]["work_location"], "WFH")

    def test_an_open_session_does_not_block_someone_else(self):
        self.clock_in(self.hanz_user)
        self.client.force_authenticate(user=self.mara_user)
        self.assertIsNone(self.client.get(reverse("records")).data["open_record"])
        self.assertEqual(self.clock_in(self.mara_user).status_code, 201)

    def test_clock_out_only_closes_your_own_session(self):
        self.clock_in(self.hanz_user)
        self.clock_in(self.mara_user)

        self.client.force_authenticate(user=self.mara_user)
        self.client.post(reverse("clock-out"))

        self.assertIsNone(AttendanceRecord.objects.get(employee=self.hanz).last_out)
        self.assertIsNotNone(AttendanceRecord.objects.get(employee=self.mara).last_out)

    def test_saving_a_shift_does_not_change_anyone_elses(self):
        self.client.force_authenticate(user=self.mara_user)
        self.client.put(reverse("profile"), {"start": "07:00", "end": "16:00"}, format="json")
        self.hanz.refresh_from_db()
        self.assertEqual(self.hanz.shift_start, time(9, 0))

    def test_export_only_contains_your_own_rows(self):
        self.clock_in(self.hanz_user)
        self.clock_in(self.mara_user, "WFH")

        self.client.force_authenticate(user=self.mara_user)
        response = self.client.get(reverse("export-excel"))
        workbook = load_workbook(filename=io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([cell.value for cell in sheet[2]][1], "WFH")
